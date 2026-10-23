"""CSV kaynakları, bilimsel etiketler ve yeniden üretilebilir Word içeriği."""

from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from gorsellestirme import grafikler
from rapor.kalite import raporVerileriniDogrula
from rapor.rapor_olustur import raporOlustur


class RaporGrafikTestleri(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.gecici = tempfile.TemporaryDirectory()
        cls.yol = Path(cls.gecici.name) / 'rapor.docx'
        cls.belge = raporOlustur(cls.yol)
        cls.metin = '\n'.join(p.text for p in cls.belge.paragraphs)

    @classmethod
    def tearDownClass(cls):
        cls.gecici.cleanup()

    def test_csv_kapsami_ve_ham_ozet_uyumu(self):
        self.assertTrue(raporVerileriniDogrula())
        self.assertEqual(len(grafikler.ozet_yukle()), 168)
        self.assertEqual(sum(map(len, grafikler.ham_yukle().values())), 840)
        for ad, adet in [('pivot', 160), ('gap', 160), ('buyuk_veri', 80)]:
            kayitlar = grafikler.oku(grafikler.SONUCLAR / 'bonuslar' / f'{ad}_ham.csv')
            self.assertEqual(len(kayitlar), adet)
            self.assertTrue(all(k['basarili_mi'] == 'True' for k in kayitlar))
        self.assertEqual(len(grafikler.oku(grafikler.SONUCLAR / 'bellek_sonuclari.csv')), 44)

    def test_bozuk_ozet_raporlanmaz(self):
        oku = grafikler.oku
        def bozuk(yol):
            satirlar = oku(yol)
            if Path(yol).name == 'ozet_100.csv':
                satirlar[0]['ortalama_sure_ms'] = '999999'
            return satirlar
        with patch.object(grafikler, 'oku', side_effect=bozuk):
            with self.assertRaisesRegex(ValueError, 'uyuşmazlığı'):
                raporVerileriniDogrula()

    def test_word_dokuz_ana_bolum(self):
        belge = Document(self.yol)
        basliklar = [p.text for p in belge.paragraphs if p.style.name == 'Heading 1']
        self.assertEqual(basliklar[:8], [
            '1. Giriş', '2. Teorik Bilgi', '3. Algoritmaların Gerçeklenmesi',
            '4. Deneysel Çalışma', '5. Performans Sonuçları',
            '6. Karşılaştırma ve Tartışma', '7. Sonuç', '8. Kaynakça'])
        self.assertTrue(any(p.style.name == 'Title' and 'Sıralama' in p.text for p in belge.paragraphs))
        self.assertEqual(len(basliklar), len(set(basliklar)))

    def test_kapak_ve_bicim(self):
        from docx.shared import Cm, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        for bilgi in ['Gülce Duru KOÇ', '22181616036', 'Doç. Dr. Adem Tekerek', '23 Ekim 2026']:
            self.assertIn(bilgi, self.metin)
        self.assertNotIn('................................', self.metin)
        normal = self.belge.styles['Normal']
        self.assertEqual(normal.font.name, 'Times New Roman')
        self.assertEqual(normal.font.size, Pt(12))
        self.assertEqual(normal.paragraph_format.line_spacing, 1.5)
        for ad, boyut in [('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 12)]:
            stil = self.belge.styles[ad]
            self.assertEqual(stil.font.size, Pt(boyut))
            self.assertTrue(stil.font.bold)
            self.assertEqual(str(stil.font.color.rgb), '000000')
        bolum = self.belge.sections[0]
        for kenar in [bolum.left_margin, bolum.right_margin, bolum.top_margin, bolum.bottom_margin]:
            self.assertAlmostEqual(kenar.cm, 2.5, places=2)
        self.assertAlmostEqual(bolum.page_width.cm, 21, places=2)
        self.assertAlmostEqual(bolum.page_height.cm, 29.7, places=2)
        self.assertEqual(bolum.footer.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertIn('PAGE', bolum.footer._element.xml)
        for p in self.belge.paragraphs:
            if re.match(r'^(Tablo|Şekil) \d+\. ', p.text):
                self.assertTrue(all(r.font.size == Pt(10) for r in p.runs))

    def test_mimari_ve_ekler(self):
        for ifade in ['3.6 Proje klasörleri', '3.7 Veri akışı', '3.8 CSV kayıtları',
                      '3.9 Otomatik test', 'dizi = baslangic_dizisi.copy()',
                      'kontrol noktalarına', 'Proje Kaynak Koduna Erişim',
                      'https://github.com/Gulced/AlgoritmaSiralamaKarsilastirma']:
            self.assertIn(ifade, self.metin)
        self.assertEqual(len(self.belge.inline_shapes), 35)

    def test_son_duzen_ekler_yok_gercek_gorseller_var(self):
        self.assertNotIn('Ek A', self.metin)
        self.assertNotIn('Ekler ve Proje Deposu', self.metin)
        self.assertNotIn('görsel doğrulaması tamamlanmadı', self.metin)
        self.assertIn('tarafımdan kontrol', self.metin)
        self.assertIn('5.13 Tüm zorunlu deney sonuçları', self.metin)
        self.assertIn('4.5.1 CSV kayıtlarından', self.metin)
        self.assertTrue(any(len(t.rows) == 169 for t in self.belge.tables))
        self.assertIn('TOC', self.belge._element.xml)
        self.assertTrue(any(p.text.startswith('- CPU sıcaklığı') for p in self.belge.paragraphs))

    def test_teslim_dil_duzeltmeleri(self):
        tablo_metni = '\n'.join(c.text for t in self.belge.tables for r in t.rows for c in r.cells)
        for yanlis in ["Şekil 10'de", "Şekil 35'de", "2.800'dır", "1,10'dir", "3.337.856'dir",
                       'zaman sınırını belirler', 'kullanıcı tarafından', 'kullanıcının',
                       'rapor_olustur.py', 'rapor/kalite.py', 'Word üreticileri']:
            self.assertNotIn(yanlis, self.metin)
        self.assertIn('çağrı yığını üst sınırı', tablo_metni)
        self.assertNotIn('cagri yigini ust siniri', tablo_metni)
        self.assertIn('timeout mekanizması içermez', self.metin)

    def test_bilimsel_ayrimlar_ve_shell_tablosu(self):
        shell = next(t for t in self.belge.tables if any('Shell (klasik)' in c.text for r in t.rows for c in r.cells))
        satir = next(r for r in shell.rows if r.cells[0].text == 'Shell (klasik)')
        self.assertEqual(satir.cells[3].text, 'O(n²)')
        self.assertEqual(satir.cells[4].text, 'O(log n)*')
        for ifade in ['gerçek çağrı', 'bölümleme ağacı', 'CPU sıcaklığı', 'kaydedilmedi',
                      'gerçek ölçüm değil', 'küçük parçayı özyinelemeye', '[50, 25, 12, 6, 3, 1]']:
            self.assertIn(ifade, self.metin)
        self.assertNotIn('aralıklar 2\'nin kuvvetleridir', self.metin)
        self.assertNotIn('O(n log n) sınıfı algoritmalar (Shell', self.metin)

    def test_grafik_turleri_ve_hareket_sifiri(self):
        grafikler.grafikOrtaminiHazirla()
        kayitlar = {}
        def yakala(fig, dosya, baslik, aciklama):
            kayitlar[dosya] = (baslik, aciklama)
            if dosya == '11_hareket_rastgele.png':
                self.assertEqual(fig.axes[0].get_yscale(), 'symlog')
                veri = grafikler.ozet_yukle()
                self.assertEqual(list(fig.axes[0].lines[0].get_ydata()),
                                 [veri['selection', 'rastgele', n]['hareket'] for n in grafikler.BOYUTLAR])
            grafikler.plt.close(fig)
        with patch.object(grafikler, 'kaydet', side_effect=yakala):
            grafikler.boyuta_gore(grafikler.ozet_yukle())
            grafikler.dagilim_etkisi(grafikler.ozet_yukle())
        for ad in ['01_sure_rastgele.png', '06_karsilastirma_rastgele.png',
                   '11_hareket_rastgele.png', '33_dagilim_sure_n100000.png']:
            self.assertIn(ad, kayitlar)

    def test_paralel_tahminler_ayri_etiketli(self):
        satirlar = grafikler.oku(grafikler.SONUCLAR / 'analizler' / 'paralel_merge_teorik.csv')
        self.assertEqual(sum(k['kayit_turu'] == 'teorik_tahmin' for k in satirlar), 12)
        self.assertEqual(sum(k['kayit_turu'] == 'gercek_sirali_olcum' for k in satirlar), 4)
        self.assertIn('paralel ölçüm', self.metin)
        grafikler.grafikOrtaminiHazirla()
        yakalanan = []
        def yakala(fig, dosya, baslik, aciklama):
            yakalanan.append((baslik, aciklama))
            grafikler.plt.close(fig)
        with patch.object(grafikler, 'kaydet', side_effect=yakala):
            grafikler.paralel_merge_bonusu()
        self.assertTrue(all('teorik' in b.lower() for b, _ in yakalanan))
        self.assertTrue(all('model' in a for _, a in yakalanan))


if __name__ == '__main__':
    unittest.main()
