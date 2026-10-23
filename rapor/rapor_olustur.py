"""BMT317 proje raporunu (Word) sonuç CSV'lerinden ve grafiklerden üretir.

Tablolar ve rapordaki türetilmiş sayılar (üs, oran vb.) doğrudan `sonuclar/`
altındaki CSV dosyalarından hesaplanır; metne elle sayı yazılan yerler yalnızca
bu çıktılara bakılarak doğrulanmıştır.

Çalıştırma (proje kökünden, python-docx ve matplotlib kurulu olmalı):
    .venv/bin/python -m rapor.rapor_olustur
Çıktı: rapor/BMT317_Siralama_Algoritmalari_Raporu.docx
"""

import math
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from rapor.kalite import raporVerileriniDogrula
from rapor.csv_gorselleri import csvKesitleriOlustur

from gorsellestirme.grafikler import (ALGORITMALAR, BOYUTLAR, DAGILIMLAR, DAGILIM_AD,
                                      SONUCLAR, oku, ozet_yukle)

KOK = Path(__file__).resolve().parents[1]
GRAFIK = SONUCLAR / "grafikler"
CIKTI = Path(__file__).resolve().parent / "BMT317_Siralama_Algoritmalari_Raporu.docx"
FONT = "Times New Roman"
ALGO_AD = {a: a.capitalize() for a in ALGORITMALAR}

def raporOlustur(cikti=CIKTI):
    raporVerileriniDogrula()
    V = ozet_yukle()  # (algoritma, dağılım, n) -> süre ms / karşılaştırma / hareket


    # ------------------------------------------------------------- sayı biçimleri
    def tr(x, ondalik=0):
        """Türkçe sayı biçimi: binlik nokta, ondalık virgül."""
        s = f"{x:,.{ondalik}f}"
        return s.replace(",", "§").replace(".", ",").replace("§", ".")


    def sure_tr(ms):
        if ms < 10:
            return tr(ms, 3)
        if ms < 1000:
            return tr(ms, 2)
        return tr(ms, 1)


    def uzun_sure(ms):
        if ms >= 60_000:
            return f"{tr(ms / 60_000, 1)} dk"
        if ms >= 1000:
            return f"{tr(ms / 1000, 1)} sn"
        return f"{tr(ms, 1)} ms"


    def us(algo, d, n1=10000, n2=100000):
        """Deneysel büyüme üssü: süre ∝ n^k."""
        return math.log(V[algo, d, n2]["sure"] / V[algo, d, n1]["sure"]) / math.log(n2 / n1)


    # ----------------------------------------------------------------- belge kurulumu
    belge = Document()
    bolum = belge.sections[0]
    bolum.page_width, bolum.page_height = Cm(21), Cm(29.7)
    bolum.left_margin = bolum.right_margin = Cm(2.5)
    bolum.top_margin = bolum.bottom_margin = Cm(2.5)
    ICERIK_GENISLIGI = 16.0  # cm

    belge.core_properties.title = "Sıralama Algoritmalarının Performans Karşılaştırması"
    belge.core_properties.subject = "BMT317 – Algoritmalar Proje Ödevi"
    belge.core_properties.author = "Gülce Duru KOÇ"


    def yazi_tipi_ayarla(stil, boyut, kalin=None, renk=RGBColor(0, 0, 0)):
        stil.font.name = FONT
        stil.font.size = Pt(boyut)
        stil.font.color.rgb = renk
        if kalin is not None:
            stil.font.bold = kalin
        rpr = stil.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for oznitelik in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(oznitelik)) is not None:
                del rfonts.attrib[qn(oznitelik)]
        for oznitelik in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(oznitelik), FONT)


    normal = belge.styles["Normal"]
    yazi_tipi_ayarla(normal, 12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5
    for ad, boyut in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        stil = belge.styles[ad]
        yazi_tipi_ayarla(stil, boyut, kalin=True)
        stil.paragraph_format.space_before = Pt(16 if ad == "Heading 1" else 12)
        stil.paragraph_format.space_after = Pt(6)
        stil.paragraph_format.keep_with_next = True
    for stil in belge.styles:
        for kenarlik in list(stil.element.iter(qn("w:pBdr"))):
            kenarlik.getparent().remove(kenarlik)
    for ad in ("List Bullet", "List Number"):
        yazi_tipi_ayarla(belge.styles[ad], 12)
        belge.styles[ad].paragraph_format.space_after = Pt(3)

    for ad in ("TOC 1", "TOC 2", "TOC 3", "Hyperlink", "FollowedHyperlink"):
        if ad not in belge.styles:
            from docx.enum.style import WD_STYLE_TYPE
            belge.styles.add_style(ad, WD_STYLE_TYPE.PARAGRAPH if ad.startswith("TOC") else WD_STYLE_TYPE.CHARACTER)
        yazi_tipi_ayarla(belge.styles[ad], 12, kalin=False)
        belge.styles[ad].font.underline = False
        if ad.startswith("TOC"):
            belge.styles[ad].paragraph_format.line_spacing = 1.5
            belge.styles[ad].paragraph_format.space_after = Pt(0)

    SAYAC = {"sekil": 0, "tablo": 0}
    sonraki_sayfa = False


    def _calistir(paragraf, metin, kalin=False, italik=False, boyut=None):
        """**kalın** işaretlerini destekleyen koşu ekleyici."""
        for parca in re.split(r"(\*\*.+?\*\*|`[^`]+`)", metin):
            if not parca:
                continue
            kalin_parca = parca.startswith("**") and parca.endswith("**")
            kod_parca = parca.startswith("`") and parca.endswith("`")
            icerik = parca[2:-2] if kalin_parca else parca[1:-1] if kod_parca else parca
            kosu = paragraf.add_run(icerik)
            kosu.font.name = FONT
            kosu.font.color.rgb = RGBColor(0, 0, 0)
            kosu.bold = kalin or kalin_parca
            kosu.italic = italik
            if boyut:
                kosu.font.size = Pt(boyut)
        return paragraf


    def P(metin, kalin=False, italik=False, hiza=WD_ALIGN_PARAGRAPH.JUSTIFY, boyut=None, sonra=None):
        nonlocal sonraki_sayfa
        paragraf = belge.add_paragraph()
        if sonraki_sayfa:
            paragraf.paragraph_format.page_break_before = True
            sonraki_sayfa = False
        paragraf.alignment = hiza
        if sonra is not None:
            paragraf.paragraph_format.space_after = Pt(sonra)
        return _calistir(paragraf, metin, kalin, italik, boyut)


    def MADDE(metin):
        paragraf = belge.add_paragraph()
        paragraf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return _calistir(paragraf, "- " + metin)


    def H(metin, seviye=1):
        nonlocal sonraki_sayfa
        paragraf = belge.add_heading(metin, level=seviye)
        if sonraki_sayfa:
            paragraf.paragraph_format.page_break_before = True
            sonraki_sayfa = False
        return paragraf


    def SAYFA_SONU():
        nonlocal sonraki_sayfa
        sonraki_sayfa = True


    def _golge(hucre, renk):
        tcpr = hucre._tc.get_or_add_tcPr()
        golge = OxmlElement("w:shd")
        golge.set(qn("w:val"), "clear")
        golge.set(qn("w:color"), "auto")
        golge.set(qn("w:fill"), renk)
        tcpr.append(golge)


    def _hucre_yaz(hucre, metin, boyut, kalin=False, hiza=WD_ALIGN_PARAGRAPH.LEFT):
        hucre.text = ""
        paragraf = hucre.paragraphs[0]
        paragraf.alignment = hiza
        paragraf.paragraph_format.space_after = Pt(0)
        paragraf.paragraph_format.line_spacing = 1.5
        kosu = paragraf.add_run(str(metin))
        kosu.font.name = FONT
        kosu.font.color.rgb = RGBColor(0, 0, 0)
        kosu.font.size = Pt(12)
        kosu.bold = kalin


    def TABLO(baslik, kolonlar, satirlar, genislikler, sagda=(), boyut=9.5, ilk_kalin=False):
        """Başlık satırı her sayfada tekrarlanan, kenarlıklı tablo ekler."""
        SAYAC["tablo"] += 1
        kapsam = belge.add_paragraph()
        kapsam.alignment = WD_ALIGN_PARAGRAPH.LEFT
        kapsam.paragraph_format.keep_with_next = True
        kapsam.paragraph_format.space_before = Pt(8)
        kapsam.paragraph_format.space_after = Pt(3)
        _calistir(kapsam, f"Tablo {SAYAC['tablo']}. {baslik}", kalin=True, boyut=10)

        toplam = sum(genislikler)
        genislikler = [g * min(1, ICERIK_GENISLIGI / toplam) for g in genislikler]
        tablo = belge.add_table(rows=1, cols=len(kolonlar))
        tablo.style = "Table Grid"
        tablo.alignment = WD_TABLE_ALIGNMENT.CENTER
        tablo.autofit = False
        cizgiler = OxmlElement("w:tblBorders")
        for kenar in ("top", "left", "bottom", "right", "insideH", "insideV"):
            cizgi = OxmlElement("w:" + kenar)
            for ad, deger in (("val", "single"), ("sz", "4"), ("color", "000000")):
                cizgi.set(qn("w:" + ad), deger)
            cizgiler.append(cizgi)
        tablo._tbl.tblPr.append(cizgiler)
        for kolon, genislik in zip(tablo.columns, genislikler):
            kolon.width = Cm(genislik)
        for i, (kolon, genislik) in enumerate(zip(kolonlar, genislikler)):
            hucre = tablo.rows[0].cells[i]
            hucre.width = Cm(genislik)
            _hucre_yaz(hucre, kolon, boyut, kalin=True,
                       hiza=WD_ALIGN_PARAGRAPH.RIGHT if i in sagda else WD_ALIGN_PARAGRAPH.LEFT)
            _golge(hucre, "FFFFFF")
        trpr = tablo.rows[0]._tr.get_or_add_trPr()
        tekrar = OxmlElement("w:tblHeader")
        tekrar.set(qn("w:val"), "true")
        trpr.append(tekrar)
        for satir in satirlar:
            hucreler = tablo.add_row().cells
            for i, (deger, genislik) in enumerate(zip(satir, genislikler)):
                hucreler[i].width = Cm(genislik)
                _hucre_yaz(hucreler[i], deger, boyut, kalin=ilk_kalin and i == 0,
                           hiza=WD_ALIGN_PARAGRAPH.RIGHT if i in sagda else WD_ALIGN_PARAGRAPH.LEFT)
            # Satırın sayfa sonunda ikiye bölünmesini engelle.
            cant = OxmlElement("w:cantSplit")
            cant.set(qn("w:val"), "true")
            tablo.rows[-1]._tr.get_or_add_trPr().append(cant)

        # Kısa tabloların tamamı aynı sayfada kalsın (son satır hariç her satır sonrakiyle birlikte).
        if len(satirlar) <= 10:
            for satir_nesnesi in tablo.rows[:-1]:
                for hucre in satir_nesnesi.cells:
                    for paragraf in hucre.paragraphs:
                        paragraf.paragraph_format.keep_with_next = True

        return tablo


    def SEKIL(dosya, baslik, genislik_cm=15.0):
        """Grafiği ve numaralı açıklamasını ekler."""
        yol = GRAFIK / dosya
        SAYAC["sekil"] += 1
        with Image.open(yol) as resim:
            oran = resim.height / resim.width
        genislik = min(genislik_cm, ICERIK_GENISLIGI)
        if genislik * oran > 16.5:
            genislik = 16.5 / oran
        paragraf = belge.add_paragraph()
        paragraf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraf.paragraph_format.keep_with_next = True
        paragraf.paragraph_format.space_before = Pt(6)
        paragraf.add_run().add_picture(str(yol), width=Cm(genislik))
        aciklama = belge.add_paragraph()
        aciklama.alignment = WD_ALIGN_PARAGRAPH.CENTER
        aciklama.paragraph_format.space_after = Pt(10)
        _calistir(aciklama, f"Şekil {SAYAC['sekil']}. {baslik}", italik=True, boyut=10)


    def alan_ekle(paragraf, talimat):
        """Word alan kodu (PAGE, TOC) ekler."""
        for tur, metin in (("begin", None), (None, talimat), ("separate", None), ("text", " "), ("end", None)):
            kosu = paragraf.add_run()
            if tur == "text":
                kosu.text = metin
                continue
            if metin is not None and tur is None:
                gosterge = OxmlElement("w:instrText")
                gosterge.set(qn("xml:space"), "preserve")
                gosterge.text = metin
                kosu._r.append(gosterge)
                continue
            alan = OxmlElement("w:fldChar")
            alan.set(qn("w:fldCharType"), tur)
            kosu._r.append(alan)


    # ----------------------------------------------------------------- alt bilgi
    sayfa_no = bolum.footer.paragraphs[0]
    sayfa_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
    alan_ekle(sayfa_no, "PAGE")
    for kosu in sayfa_no.runs:
        kosu.font.name = FONT
        kosu.font.size = Pt(10)

    # ====================================================================== KAPAK
    for metin in ("GAZİ ÜNİVERSİTESİ", "TEKNOLOJİ FAKÜLTESİ",
                  "BİLGİSAYAR MÜHENDİSLİĞİ BÖLÜMÜ"):
        P(metin, kalin=True, hiza=WD_ALIGN_PARAGRAPH.CENTER, boyut=14, sonra=6)
    P("BMT317 – Algoritmalar", kalin=True, hiza=WD_ALIGN_PARAGRAPH.CENTER, boyut=14, sonra=18)
    baslik = P("Sıralama Algoritmalarının Performans Karşılaştırması", kalin=True,
               hiza=WD_ALIGN_PARAGRAPH.CENTER, boyut=14, sonra=24)
    baslik.style = belge.styles["Title"]
    yazi_tipi_ayarla(belge.styles["Title"], 14, kalin=True)
    P("Öğrenci Adı Soyadı: Gülce Duru KOÇ", hiza=WD_ALIGN_PARAGRAPH.CENTER)
    P("Öğrenci Numarası: 22181616036", hiza=WD_ALIGN_PARAGRAPH.CENTER)
    P("Öğretim Elemanı: Doç. Dr. Adem Tekerek", hiza=WD_ALIGN_PARAGRAPH.CENTER)
    P("Tarih: 23 Ekim 2026", hiza=WD_ALIGN_PARAGRAPH.CENTER)

    SAYFA_SONU()

    # ================================================================ İÇİNDEKİLER
    P("İçindekiler", kalin=True, boyut=14, hiza=WD_ALIGN_PARAGRAPH.LEFT)
    # Word, başlık stillerinden sayfa numaralı içindekileri hesaplar.
    icindekiler = belge.add_paragraph()
    alan_ekle(icindekiler, 'TOC \\o "1-2" \\h \\z \\u')
    SAYFA_SONU()

    # ===================================================================== 1. GİRİŞ
    H("1. Giriş")
    P("Bu projede yedi sıralama algoritmasını aynı veri üretimi ve sayaç tanımlarıyla karşılaştırdım. "
      "Çalışmanın odağı, yalnızca bir algoritmanın daha hızlı olduğunu göstermek değil; süre farkının "
      "hangi karşılaştırma ve veri hareketlerinden kaynaklandığını incelemektir. Örneğin 100.000 "
      "elemanlı rastgele veride Selection Sort'un karşılaştırma sayısı n(n−1)/2 olurken Merge Sort "
      "diziyi seviyeler hâlinde birleştirir. Bu farkı gerçek CSV kayıtlarındaki süre ve sayaçlarla "
      "birlikte değerlendirdim. Bellek ölçümleri süre deneylerinden ayrı yürütüldü.")

    P("**Temel araştırma sorusu.** Farklı veri büyüklükleri ve veri dağılımları için yedi algoritmadan "
      "hangisi veya hangileri daha avantajlıdır ve elde edilen sonuçlar algoritmaların teorik zaman ve "
      "bellek karmaşıklıklarıyla ne ölçüde örtüşmektedir?")
    P("**Kapsam.** Zorunlu deney matrisi 6 veri boyutu (100; 1.000; 5.000; 10.000; 50.000; 100.000), 4 veri "
      "dağılımı (rastgele, sıralı, ters sıralı, kısmen sıralı) ve 7 algoritmadan oluşan 168 kombinasyondur. "
      "Her kombinasyon bir ısınma (warm-up) ve beş gerçek ölçümle çalıştırılmış, toplam **840 ölçüm** "
      "yapılmıştır; tüm ölçümlerde sonuç dizisi doğrulanmıştır. Ayrıca bellek analizi ve ödevin beş "
      "bonus çalışması yapılmıştır: Quick Sort pivot karşılaştırması, Shell Sort gap karşılaştırması, "
      "n = 500.000 deneyi, paralel Merge Sort'un teorik incelemesi ve Bubble Sort adım görselleştirmesi.")
    P("**Raporun yapısı.** Bölüm 2 teorik bilgiyi, Bölüm 3 gerçeklemeleri, Bölüm 4 deney tasarımını ve "
      "ortamı, Bölüm 5 ölçüm sonuçlarını sunar. Bölüm 6 sonuçları teoriyle karşılaştırır ve ödevde "
      "sorulan soruları yanıtlar; Bölüm 7 genel sonuçları verir.")

    # ============================================================ 2. TEORİK BİLGİ
    H("2. Teorik Bilgi")
    H("2.1 Karşılaştırma ölçütleri", 2)
    P("Bir sıralama algoritmasının maliyeti genellikle üç nicelikle anlatılır: veri karşılaştırma sayısı, "
      "veri hareketi (atama ya da yer değiştirme) sayısı ve ek bellek gereksinimi. Karşılaştırma tabanlı "
      "sıralamada en kötü durumda en az Ω(n log n) karşılaştırma gerektiği bilinir; bu yüzden n log n "
      "sınıfı bu problem için asimptotik olarak en iyi sınıftır. Big-O gösterimi sabit çarpanları ve düşük "
      "mertebeden terimleri gizler; bu çalışma, gizlenen bu ayrıntıların gerçek süreye etkisini de "
      "gözlemlemeyi amaçlar (Cormen ve diğerleri, 2022, sıralama bölümleri).")

    H("2.2 Algoritmalar ve teorik karmaşıklıklar", 2)
    TABLO("Algoritmaların teorik karmaşıklıkları",
          ["Algoritma", "En iyi", "Ortalama", "En kötü", "Ek bellek", "Kararlı"],
          [["Selection", "O(n²)", "O(n²)", "O(n²)", "O(1)", "Hayır"],
           ["Insertion", "O(n)", "O(n²)", "O(n²)", "O(1)", "Evet"],
           ["Bubble (erken durmalı)", "O(n)", "O(n²)", "O(n²)", "O(1)", "Evet"],
           ["Shell (klasik)", "O(n log n)", "Gap/veriye bağlı", "O(n²)", "O(log n)*", "Hayır"],
           ["Quick", "O(n log n)", "O(n log n)", "O(n²)", "O(log n)", "Hayır"],
           ["Merge", "O(n log n)", "O(n log n)", "O(n log n)", "O(n)", "Evet"],
           ["Heap", "O(n)*", "O(n log n)", "O(n log n)", "O(1)", "Hayır"]],
          [3.6, 2.3, 2.8, 2.3, 2.4, 2.2], ilk_kalin=True)
    P("* Bu çalışmadaki Shell Sort gerçeklemesi gap listesini bir listede sakladığı için O(log n) ek alan "
      "kullanır. Quick Sort'un O(log n) değeri, küçük parçaya özyineleme yapan gerçeklemeye aittir; "
      "bu önlem alınmasa özyineleme derinliği en kötü durumda O(n) olurdu.", italik=True, boyut=12)

    P("Heap için O(n) en iyi durum tüm değerler eşit olduğunda mümkündür: aşağı süzme ilk "
      "seviyede durur. Deney verileri farklı 0..n−1 değerleri içerdiğinden bu özel durum ana "
      "deneylerde yer almaz. Teorik tabloda Quick'in ortalaması rastgele girdi varsayımına dayanır.",
      italik=True, boyut=12)

    H("2.3 Shell Sort ve gap dizileri", 2)
    P("Projede klasik gap listesi önce n//2, ardından gap//2 hesaplanarak oluşturulur. n=10 için "
      "[5, 2, 1], n=100 için [50, 25, 12, 6, 3, 1] elde edilir; dolayısıyla değerler her zaman "
      "ikinin kuvvetleri değildir. Her gap için kaydırmalı Insertion Sort uygulanır. Bu klasik "
      "yöntemin genel en kötü durum üst sınırı O(n²)'dir (Knuth, 1998, Bölüm 5.2.1). Ölçümlerde "
      "daha yavaş bir büyüme görülmesi bu sınırı değiştirmez. Bonus yönteminde h=3h+1 ile üretilen "
      "ve n'den küçük olan Knuth aralıkları ters sırayla kullanılır. Knuth dizisinin klasik "
      "incelemesi O(n√n) en kötü durum sınırı verir (Sedgewick ve Wayne, 2011; Shell sınıfı "
      "belgeleri). Bizim kodumuz her iki yöntemde de gap listesini sakladığından yardımcı alan O(log n)'dir.")

    H("2.4 Quick Sort ve pivot seçimi", 2)
    P("Quick Sort'ta orta pivot, alt dizinin (başlangıç+bitiş)//2 indeksindeki değerdir; medyan "
      "hesaplanmaz. İki uçtan ilerleyen indeksler pivotun yanlış tarafında kalan değerleri değiştirir. "
      "Dengeli bölünmelerde bölümleme ağacının yüksekliği O(log n), toplam iş O(n log n) olur. "
      "Tekrarlanan dengesiz bölünmelerde ağaç yüksekliği O(n), süre O(n²) olabilir (Hoare, 1962; "
      "Cormen ve diğerleri, 2022). Bu ağacın yüksekliği, projedeki Python çağrı yığınıyla aynı değildir. "
      "`_araligi_sirala` küçük parçayı özyinelemeye verir, büyük parçanın sınırlarını değiştirip döngüye "
      "devam eder. İç içe her çağrı en fazla yarı büyüklükte bir parça aldığı için gerçek çağrı "
      "yığını O(log n) ile sınırlıdır; bu önlem karesel karşılaştırma maliyetini ortadan kaldırmaz.")

    H("2.5 Bellek kullanımı", 2)
    P("Merge Sort birleştirme için n elemanlı yardımcı dizi gerektirir (O(n)). Quick Sort yerinde "
      "çalışır ancak özyineleme yığını kullanır; en kötü durumda O(n) olabilecek yığın derinliği, "
      "önce küçük parçanın özyinelemeyle sıralanmasıyla O(log n) ile sınırlanır. Heap Sort tamamen "
      "yerinde çalışır ve yinelemeli süzme kullanıldığında O(1) ek bellek yeterlidir. Selection, "
      "Insertion ve Bubble Sort da O(1) ek bellek kullanır.")

    # ================================================ 3. ALGORİTMALARIN GERÇEKLENMESİ
    H("3. Algoritmaların Gerçeklenmesi")
    H("3.1 Genel yapı", 2)
    P("Yedi algoritma Python ile, hazır `sort()` veya `sorted()` çağrılarına dayanmadan, "
      "`algoritmalar/` klasöründe ayrı modüller olarak yazılmıştır. Her fonksiyon "
      "`fonksiyon(dizi, sayaclar)` biçiminde çağrılır, verilen listeyi yerinde sıralar ve "
      "karşılaştırma/hareket sayılarını `Sayaclar` nesnesine ekler. Algoritmalar deney, kayıt ve arayüz "
      "kodundan bağımsızdır; böylece aynı kod hem testlerde hem ölçümlerde kullanılır.")

    H("3.1.1 Neden Python?", 3)
    P("Bu çalışmada amaç, sıralama algoritmalarının adımlarını açık biçimde göstermek ve aynı "
      "ölçüm kuralları altında karşılaştırmaktır. Python'un okunabilir sözdizimi; karşılaştırma, "
      "kaydırma ve takas işlemlerini kod üzerinde izlemeyi kolaylaştırır. Algoritmaları aynı "
      "liste yapısı ve ortak `Sayaclar` nesnesiyle uygulamak, işlem sayılarının tutarlı "
      "tanımlanmasını ve doğruluk testlerinin ortak bir yapıda yürütülmesini sağlar.")
    P("Standart kütüphanedeki `unittest`, `random.Random`, `perf_counter_ns`, `csv` ve "
      "`tracemalloc` araçları test, tekrar üretilebilir veri, süre ölçümü, kayıt ve bellek "
      "analizi için yeterlidir. Bu nedenle algoritma ve deney altyapısı ek paket gerektirmez. "
      "Buna karşılık ölçülen süreler Python yorumlayıcısının ve sayaç artırımlarının maliyetini "
      "de içerir. Sonuçları diller arasında bir hız karşılaştırması olarak değil, bu projedeki "
      "Python gerçeklemelerinin karşılaştırması olarak değerlendiriyorum.")

    H("3.2 Algoritma ayrıntıları", 2)
    MADDE("**Selection Sort:** Sıralanmamış bölümün en küçük elemanını bulur ve başa yerleştirir. "
          "Konumu zaten doğruysa yer değiştirme yapmaz; bu nedenle hareket sayısı çok düşüktür.")
    MADDE("**Insertion Sort:** Her elemanı geçiciye alıp kendisinden büyükleri sağa kaydırarak yerine "
          "yerleştirir. Başarısız son karşılaştırma da sayılır.")
    MADDE("**Bubble Sort:** Komşu elemanları karşılaştırıp takas eder. İki iyileştirme içerir: her turda "
          "sona yerleşen eleman dışarıda bırakılır ve takas olmayan turda algoritma erken durur.")
    MADDE("**Shell Sort:** Her aralıkta kaydırmalı eklemeli sıralama yapar. Varsayılan aralıklar n//2, ardından gap//2 ile üretilir; son aralık 1 "
          "(klasik); bonus karşılaştırmada Knuth dizisi kullanılır.")
    MADDE("**Quick Sort:** İki yönlü (Hoare tipi) bölümleme kullanır. Pivot varsayılan olarak orta "
          "indeksteki elemandır; bonus için ilk eleman seçeneği vardır. Küçük parça özyinelemeyle, büyük "
          "parça döngüyle işlenir.")
    MADDE("**Merge Sort:** Yukarıdan aşağıya (top-down) bölme uygular. Alt diziler oluşturmaz; tüm çağrılar "
          "n elemanlı tek bir yardımcı tamponu paylaşır. Eşitlikte soldan seçerek kararlılığı korur.")
    MADDE("**Heap Sort:** Max-heap kurar, kökü sona taşıyıp yığını küçülterek yinelemeli (özyinelemesiz) "
          "aşağı süzme uygular.")

    H("3.3 Sayaç kuralları", 2)
    P("Sayaçlar, algoritmanın veri üzerinde yaptığı işlemleri şu kurallarla sayar:")
    MADDE("**Karşılaştırma:** İki veri elemanının her karşılaştırması, başarısız son karşılaştırmalar "
          "dahil, 1 sayılır. İndeks ve döngü koşulları sayılmaz.")
    MADDE("**Veri hareketi:** Veri atamaları sayılır. Geçiciye eleman alma ve her dizi yazımı 1 harekettir; "
          "dolayısıyla geçici değişkenli bir swap **3 hareket** eder. Merge Sort'ta tampona ve ana diziye "
          "yazımlar, Quick Sort'ta pivot kopyası sayılır. İndeks atamaları sayılmaz.")
    P("Hareket sayısı bu nedenle “swap sayısı”ndan farklıdır: Selection ve Heap gibi swap yapan "
      "algoritmalarda hareket = 3 × swap'tır, Insertion ve Shell gibi kaydıran algoritmalarda ise her "
      "kaydırma 1 harekettir. Algoritmalar arası karşılaştırma bu ortak ölçüt üzerinden yapılmıştır.")

    H("3.4 Doğruluk kontrolü", 2)
    P("Her ölçümde projedeki `isSorted(dizi)` fonksiyonu çalıştırılmış ve ayrıca eleman "
      "bütünlüğü (çıktının girdiyle aynı elemanlardan oluşması) frekans farklarıyla denetlenmiştir; "
      "yalnızca ikisini birden geçen ölçüm başarılı sayılmıştır. Otomatik testler algoritmaları, veri üretimini, deney altyapısını ve doğrulamayı kapsar.  Zorunlu deneylerin 840 ölçümünün "
      "hepsi başarılıdır.")

    H("3.5 Adil karşılaştırma", 2)
    P("Her ölçüm başlangıç dizisinin ayrı bir kopyası üzerinde çalışır; bir algoritmanın sıraladığı dizi "
      "başka bir algoritmaya verilmez. Sayaç nesnesi her ölçümde yenidir.")

    H("3.6 Proje klasörleri ve modüllerin görevleri", 2)
    P("Aşağıdaki ağaç, rapor hazırlanırken projede bulunan dosyaların görev odaklı özetidir. "
      "Paketlerin __init__.py dosyaları ve sonuç klasörlerindeki tekil kayıtlar okunabilirlik için gösterilmemiştir.")
    agac = (
        "SiralamaKarsilastirma/", "  main.py",
        "  algoritmalar/  selection.py, insertion.py, bubble.py, shell.py, quick.py, merge.py, heap.py",
        "  cekirdek/  sayaclar.py, dogrulama.py, veri_uretimi.py",
        "  deneyler/  olcum.py, deney_yoneticisi.py, kayit.py, tam_deney.py",
        "    on_analiz.py, bellek_analizi.py, bonus_pivot.py, bonus_gap.py, bonus_buyuk_veri.py",
        "  analizler/  paralel_merge_analizi.py",
        "  gorsellestirme/  bubble_gorsel.py, grafikler.py",
        "  testler/", "  sonuclar/  deneyler/, bonuslar/, analizler/, grafikler/",
        "  README.md, requirements.txt, .gitignore",
    )
    for satir in agac:
        P(satir, hiza=WD_ALIGN_PARAGRAPH.LEFT, sonra=0)
    P("main.py küçük örnek çalıştırmayı sunar. Algoritma dosyaları yalnızca verilen dizi ve sayaçla "
      "çalışır. Çekirdek paketindeki veri üretimi yerel random.Random nesnesi kullanır; isSorted "
      "komşu değerlerin sırasını denetler. Sayaclar veri sınıfının iki alanı sıfırdan başlar ve "
      "ölçümler arasında ortak, değişebilir bir global sayaç tutulmaz.")
    P("Deney paketinde olcum.py tek algoritma çağrısının başlangıç ve bitiş zamanlarını ölçer; deney_yoneticisi.py "
      "warm-up, tekrar ve seed düzenini kurar. kayit.py CSV alanlarını ve sayı biçimlerini tanımlar. "
      "tam_deney.py seçilen grupları çalıştırır ve kontrol noktalarını yönetir. Bonus ölçümleri ayrı "
      "dosyalara yazılır. Bellek analizi süre deneylerinden ayrıdır; paralel Merge modülü ise gerçek "
      "paralel sıralama çalıştırmadan model hesabı yapar. Grafikler kayıtlı sonuçlardan hazırlanır.")

    H("3.7 Veri akışı ve ölçüm sınırı", 2)
    P("Veri akışı şu sırayı izler: veriUret → başlangıç dizisi → tekOlcum içindeki bağımsız kopya ve "
      "Sayaclar → algoritma → sıralılık ve eleman bütünlüğü → ham kayıt → beş tekrarın özeti → CSV. "
      "Veri üretici 0 ile n−1 arasındaki aynı değerleri kullanır. Kısmen sıralı veride seçilen konumlar "
      "döngüsel kaydırılır; dolayısıyla veri türleri arasında eleman kümesi değişmez.")
    P("Aşağıdaki kısa parça deneyler/olcum.py içindeki gerçek ölçüm sınırını gösterir. Kopya ve sayaç "
      "saat başlamadan hazırlanır; doğrulama bitiş zamanından sonra yapılır.")
    kaynak = (KOK / "deneyler/olcum.py").read_text(encoding="utf-8")
    bas = kaynak.index("    dizi = baslangic_dizisi.copy()")
    son = kaynak.index("\n", kaynak.index("    bitis = time.perf_counter_ns()", bas))
    for satir in kaynak[bas:son].splitlines():
        P(satir, hiza=WD_ALIGN_PARAGRAPH.LEFT, sonra=0)
    P("Eleman bütünlüğü için başlangıç dizisinin Counter frekanslarından çıktı frekansları çıkarılır. "
      "Bütün farkların sıfır olması ve isSorted sonucunun True olması gerekir. Böylece yanlışlıkla "
      "eleman silen fakat sıralı görünen bir çıktı başarılı kabul edilmez. Gerçek tekrarlar 317–321 "
      "seed değerlerini kullanır; warm-up ortalamaya alınmaz. Deney yöneticisinin çoklu algoritma "
      "çağrısında algoritma sırası tekrar indeksine göre döndürülür. Aşamalı zorunlu çalıştırıcı "
      "grupları tek algoritmayla çağırdığı için bu çağrılarda algoritmalar arası sıra dönüşümü oluşmaz.")

    H("3.8 CSV kayıtları ve kontrol noktaları", 2)
    P("Ham CSV her tekrarın seed, süre, karşılaştırma, hareket ve doğrulama sonuçlarını saklar. "
      "Özet CSV, aynı algoritma, dağılım ve boyut için beş tekrarın ortalamasını verir. Ham kayıt "
      "sayıları ve özet ortalamaları birbiriyle karşılaştırılarak denetlenir.")
    P("Tamamlanan deney grupları kontrol noktalarına kaydedilir. Yeniden başlatıldığında beş "
      "başarılı tekrarı bulunan gruplar atlanır; eksik gruplar tamamlanmış sayılmaz. Geçici "
      "dosya üzerinden kayıt, kesinti sırasında tamamlanmış grupların korunmasını sağlar. "
      "Altyapı gerçek bir ölçüm timeout mekanizması içermez; süre ölçümü işlemi durdurmaz.")

    H("3.9 Otomatik test yaklaşımı", 2)
    P("unittest testleri boş ve tek elemanlı girdilerden negatif ve tekrar eden değerlere kadar "
      "sıralama doğruluğunu sınar. Küçük örneklerde karşılaştırma ve hareket sayıları elle beklenen "
      "değerlerle karşılaştırılır; Merge için kararlılık, Quick için büyük sıralı girdiler ve "
      "geçersiz stratejiler ayrıca ele alınır. Veri üretiminde seed tekrar üretilebilirliği ve "
      "dağılımlar arasındaki eleman bütünlüğü denetlenir.")
    P("Deney altyapısı testlerinde geçici klasörler ve sahte ölçümler kullanılarak CSV başlıkları, "
      "tekrar sayıları, kesinti sonrası devam ve mükerrer kayıt önleme kontrol edilir. Bu testler "
      "uzun performans deneylerinin yerine geçmez. Tkinter testleri pencere açmadan adım kayıtlarını "
      "ve gerçek Bubble Sort ile aynı nihai sonucu denetler; bu, pencerenin masaüstünde görünmesini "
      "doğrulayan bir test değildir.")

    # ============================================================ 4. DENEYSEL ÇALIŞMA
    H("4. Deneysel Çalışma")
    H("4.1 Deney ortamı", 2)
    TABLO("Deney ortamı", ["Bileşen", "Değer"],
          [["İşlemci", "Apple M5 (10 çekirdek)"],
           ["RAM", "16 GB"],
           ["İşletim sistemi", "macOS 26.6"],
           ["Programlama dili ve sürümü", "Python 3.13.13 (CPython)"],
           ["Derleyici / çalışma ortamı", "CPython bayt kodu yorumlayıcısı; sanal ortam (.venv)"],
           ["IDE", "PyCharm"],
           ["Zamanlayıcı", "time.perf_counter_ns() (Python Software Foundation, time belgeleri)"],
           ["Bellek ölçümü", "tracemalloc"],
           ["Grafik kütüphanesi", "matplotlib"],
           ["Shell Sort gap yöntemi", "Klasik n//2, ardından gap//2 (ana deney); Knuth 1, 4, 13, … (bonus)"],
           ["Quick Sort pivot yöntemi", "Orta indeksteki eleman (ana deney); ilk eleman (bonus)"]],
          [5.5, 10.5], ilk_kalin=True, boyut=10)
    P("Ana deneylerde Shell için klasik gap ve Quick için orta indeks pivotu kullanılır. "
      "Knuth dizisi bonus deneyde gap seçiminin etkisini ölçmek için kullanılmıştır. "
      "Orta indeks pivotu sıralı ve ters sıralı girdilerde dengeli bölünme sağlar. İlk pivot "
      "bu girdilerde dengesiz bölünmenin etkisini incelemek için bonusta kullanılmıştır.")

    P("Tablodaki işlemci, RAM, çekirdek sayısı ve işletim sistemi 23.10.2026 tarihli yerel sistem "
      "okumasıyla doğrulanmıştır. Önceki CSV'lerde deney anına ait donanım ve işletim sistemi "
      "metaverisi bulunmadığından aynı yapılandırmanın her oturumda geçerli olduğu bağımsız olarak "
      "kanıtlanamaz. IDE bilgisi proje ortamında belirtilen PyCharm'dır; sürümü kaydedilmemiştir.")

    H("4.2 Veri üretimi", 2)
    P("Tüm diziler 0'dan n−1'e kadar farklı tam sayılardan oluşan permütasyonlardır; dört dağılım aynı "
      "değer kümesini içerir. Yerel bir `random.Random` üreteci kullanıldığından üretim deterministiktir ve "
      "global rastgele durum etkilenmez.")
    MADDE("**Rastgele:** Fisher–Yates karıştırması.")
    MADDE("**Sıralı:** 0, 1, …, n−1.")
    MADDE("**Ters sıralı:** n−1, n−2, …, 0.")
    MADDE("**Kısmen sıralı:** Sıralı diziden rastgele seçilen yaklaşık %10 konum (en az 2) kendi aralarında "
          "döngüsel kaydırılır; geri kalan %90 yerinde kalır.")

    H("4.3 Deney tasarımı", 2)
    P("Her (boyut, dağılım, algoritma) kombinasyonu bir ısınma ölçümü ve beş gerçek ölçümle çalıştırılmış, "
      "özet değerler beş ölçümün ortalamasıdır. Isınma ölçümü ilk çalıştırmanın başlangıç maliyetlerini "
      "azaltmak amacıyla yapılır; bu etkilerin giderildiği ayrıca ölçülmemiştir. Ortalamaya katılmaz. Isınmanın seed'i 316, "
      "gerçek tekrarların seed'leri 317–321'dir; böylece her tekrar farklı rastgele veri kullanır ve "
      "her algoritma aynı seed için aynı veriyi alır. Sıralı ve ters sıralı dağılımlar seed'den bağımsızdır, "
      "bu yüzden bu dağılımlarda beş tekrar aynı veriyle yapılır. Toplam ölçüm sayısı 6 × 4 × 7 × 5 = 840'tır.")
    P("Süre yalnızca algoritma çağrısını kapsar; veri üretimi, kopyalama, sayaç oluşturma, doğrulama ve "
      "kayıt süre dışındadır. Sayaç artırımları algoritmanın parçasıdır ve süreye dahildir.")

    H("4.4 Bellek ölçümü", 2)
    P("Ek bellek `tracemalloc` ile algoritma çağrısı sırasında oluşan tepe Python tahsisi olarak ölçülmüştür. "
      "Süre ve bellek ölçümleri ayrı çalıştırılmıştır, çünkü `tracemalloc` yorumlayıcıyı yavaşlatır. "
      "Bu ölçüm toplam RAM kullanımı ya da matematiksel yardımcı alanın kesin karşılığı değildir; "
      "izlenebilen Python tahsislerini gösterir (Python Software Foundation, tracemalloc belgeleri).")

    H("4.5 Kayıt ve teslim edilen veriler", 2)
    P("Her ham ölçüm algoritma, veri tipi, boyut, seed, tekrar numarası, süre (ns ve ms), karşılaştırma sayısı, "
      "hareket sayısı ve doğrulama sonuçlarıyla `sonuclar/deneyler/ham_<n>.csv` dosyalarına; beş tekrarın "
      "ortalaması `ozet_<n>.csv` dosyalarına yazılmıştır. Bonus ve bellek sonuçları `sonuclar/bonuslar/` "
      "ve `sonuclar/bellek_sonuclari.csv` içindedir. Raporun tablo ve grafikleri bu dosyalardan "
      "üretilmiştir (`gorsellestirme/grafikler.py`). Deneyleri yeniden çalıştırma adımları README "
      "dosyasındadır.")

    H("4.5.1 CSV kayıtlarından örnek kesitler", 3)
    P(f"Şekil {SAYAC['sekil'] + 1}–{SAYAC['sekil'] + 3}, kayıt düzenini göstermek için "
      "kaynak CSV'lerden seçilmiş üçer satır ve ilgili sütunları içerir. Dosyaların tamamı "
      "gösterilmemiştir. Sütun başlıkları okunabilirlik için kısaltılmış; hücre değerleri "
      "kaynak dosyalardan değiştirilmeden alınmıştır.")
    kesitler = csvKesitleriOlustur(KOK / "rapor" / "gorseller")
    SEKIL(kesitler[0], "ham_100000.csv kesiti: ham sonuçlar her tekrarı ayrı seed, süre ve sayaç değerleriyle kaydeder.", 16)
    SEKIL(kesitler[1], "ozet_100000.csv kesiti: özet sonuçlar beş gerçek tekrarın ortalama süre, karşılaştırma ve hareket değerlerini verir.", 16)
    SEKIL(kesitler[2], "bellek_sonuclari.csv kesiti: ayrı tracemalloc kayıtları bellek analizini destekler; süre deneylerinden bağımsızdır.", 16)

    H("4.6 Sınırlılıklar", 2)
    MADDE("CPU sıcaklığı, frekans değişimi ve arka plan yükü deney sırasında kaydedilmedi. Isınma veya "
          "işlemci hızının düşmesi olası etkenlerdir; herhangi bir süre farkının kesin nedeni olarak gösterilemez.")
    MADDE("Algoritma süreleri Python yorumlayıcısı, nesne yönetimi ve sayaç artırımı maliyetlerini "
          "birlikte içerir. C/C++ ile karşılaştırmalı deney yapılmadığından bu dillere göre hız "
          "oranı verilemez. Buradaki süreler incelenen Python gerçeklemelerine aittir.")

    MADDE("Süreler donanıma, işlemci zamanlamasına ve oturum koşullarına bağlıdır; deneyler farklı "
          "oturumlarda alındığından mutlak süreler evrensel bir performans garantisi değildir.")
    MADDE("O(n²) algoritmalar n = 100.000'de dakikalar sürmüştür (Bubble Sort rastgele veride yaklaşık "
          f"{uzun_sure(V['bubble','rastgele',100000]['sure'])}); bu ölçümler zaman aşımı olmadan tamamlanmıştır. "
          "n = 500.000 deneyine bu nedenle yalnızca önceki ölçümlerde daha kısa süren algoritmalar (Shell, Quick, Merge, "
          "Heap) alınmıştır.")
    MADDE("Bellek ölçümü yalnızca n = 1.000 ve n = 10.000 için yapılmıştır; n = 10.000'de yalnızca "
          "Shell, Quick, Merge ve Heap ölçülmüştür. Pivot ve gap bonusları n ≤ 10.000 ile sınırlıdır; "
          "ilk eleman pivotu sıralı veride O(n²)'ye düştüğü için daha büyük boyut denenmemiştir.")
    MADDE("Her ölçüm grubunda yalnızca beş tekrar vardır; bu, değişkenlik "
          "hakkında sınırlı bilgi verir; istatistiksel anlamlılık testi uygulanmamıştır.")

    # ========================================================= 5. PERFORMANS SONUÇLARI
    H("5. Performans Sonuçları")
    H("5.1 Doğruluk", 2)
    P("Zorunlu deneylerin 840 ölçümünün tamamı `isSorted` ve eleman bütünlüğü denetimini geçmiştir; "
      "168 özet kombinasyonunun tamamı başarılıdır. Bellek analizindeki 44 ölçüm, Quick Sort pivot "
      "bonusundaki 160, Shell gap bonusundaki 160 ve n = 500.000 deneyindeki 80 ölçüm de başarılıdır.")

    H("5.2 Sonuç tabloları", 2)
    P("Tüm zorunlu sonuçlar Bölüm 5.13'teki 168 satırlık tabloda verilmiştir. Aşağıda "
      "n = 1.000 ve n = 100.000 için ortalama sonuçlar yer alır; süreler milisaniyedir.")
    for n in (1000, 100000):
        satirlar = []
        for d in DAGILIMLAR:
            for a in ALGORITMALAR:
                r = V[a, d, n]
                satirlar.append([ALGO_AD[a], DAGILIM_AD[d], sure_tr(r["sure"]),
                                 tr(r["karsilastirma"]), tr(r["hareket"])])
        TABLO(f"n = {tr(n)} için ortalama sonuçlar (5 tekrar)",
              ["Algoritma", "Veri tipi", "Süre (ms)", "Karşılaştırma", "Hareket"],
              satirlar, [2.4, 3.4, 2.8, 3.7, 3.7], sagda=(2, 3, 4), boyut=9)

    H("5.3 Veri boyutuna göre çalışma süresi", 2)
    P("Aşağıdaki şekillerde süre veri boyutuna göre log-log eksende gösterilmiştir. Gri noktalı çizgi n², kesikli "
      "çizgi n·log n eğilimini temsil eder. Log-log eksende O(n²) algoritmalar eğimi yaklaşık 2, "
      "O(n log n) algoritmalar yaklaşık 1 olan doğrulara yakındır.")
    SEKIL("01_sure_rastgele.png", "Rastgele veride çalışma süresinin veri boyutuna göre değişimi (log-log).", 13)
    SEKIL("05_sure_panel.png", "Dört veri dağılımında çalışma süresi (log-log).", 14.5)
    SEKIL("20_hizli_algoritmalar_dogrusal.png",
          "Shell, Quick, Merge ve Heap, doğrusal eksen; Shell için n log n garantisi yoktur.", 15.5)
    SEKIL("21_n100000_sure_cubuk.png", "n = 100.000 rastgele veride ortalama süre (log eksen).", 12.5)

    P(f"Şekil {SAYAC['sekil']-3}–{SAYAC['sekil']} süreyi farklı ölçeklerde gösterir. Rastgele veride "
      f"100.000 elemanda Quick {uzun_sure(V['quick','rastgele',100000]['sure'])}, Selection "
      f"{uzun_sure(V['selection','rastgele',100000]['sure'])} sürer. Quick'in bölümlemesi ile "
      "Selection'ın tüm kalan bölümü taraması arasındaki iş farkı bu ayrımı açıklar. Referans "
      "çizgileri yalnızca ölçeklenmiş büyüme eğilimleridir, ölçüm veya performans garantisi değildir.")
    H("5.4 Karşılaştırma sayısı", 2)
    SEKIL("10_karsilastirma_panel.png", "Dört veri dağılımında ortalama karşılaştırma sayısı (log-log).", 14.5)
    SEKIL("52_eleman_basi_karsilastirma.png",
          "Shell, Quick, Merge ve Heap için eleman başına karşılaştırma; log₂ n referanstır.", 12)

    P(f"Şekil {SAYAC['sekil']-1}'te rastgele 100.000 eleman için Selection "
      f"{tr(V['selection','rastgele',100000]['karsilastirma'])}, Merge "
      f"{tr(V['merge','rastgele',100000]['karsilastirma'])} karşılaştırma yapar. Selection'ın "
      "n(n−1)/2 tarama sayısı dağılımdan bağımsızdır; Merge'de birleştirme sırasında yarılardan "
      "birinin erken bitmesi karşılaştırma sayısını değiştirir.")
    H("5.5 Swap / veri hareketi", 2)
    SEKIL("15_hareket_panel.png",
          "Dört veri dağılımında ortalama veri hareketi sayısı (symlog eksen; 0 değerleri gösterilebilir).", 13.5)
    SEKIL("53_eleman_basi_hareket.png", "Shell, Quick, Merge ve Heap için eleman başına veri hareketi.", 12)

    P(f"Şekil {SAYAC['sekil']-1}'da 100.000 elemanlı sıralı veride Bubble "
      f"{tr(V['bubble','sirali',100000]['hareket'])}, Insertion "
      f"{tr(V['insertion','sirali',100000]['hareket'])} hareket yapar. Bubble takas yapmadan durur; "
      "Insertion ise her elemanı geçiciye alıp tekrar yazar. Symlog ölçeği sıfır hareketi de gösterir. "
      f"Şekil {SAYAC['sekil']} rastgele girdiye aittir: Quick eleman başına "
      f"{tr(V['quick','rastgele',100000]['hareket']/100000,2)}, Merge "
      f"{tr(V['merge','rastgele',100000]['hareket']/100000,2)} veri ataması yapar.")
    H("5.6 Veri dağılımının performansa etkisi", 2)
    SEKIL("33_dagilim_sure_n100000.png", "Dağılımın çalışma süresine etkisi, n = 100.000 (log eksen).", 14)
    SEKIL("34_dagilim_karsilastirma_n100000.png", "Dağılımın karşılaştırma sayısına etkisi, n = 100.000.", 14)
    SEKIL("35_dagilim_hareket_n100000.png", "Dağılımın veri hareketine etkisi, n = 100.000.", 14)
    SEKIL("37_dagilim_oran_n100000.png", "Rastgele veriye göre süre oranı, n = 100.000.", 14)
    SEKIL("40_isi_sure_n100000.png", "Algoritma × dağılım ısı haritası (süre), n = 100.000.", 13)
    SEKIL("56_dagilim_duyarliligi.png", "Algoritmaların dağılım duyarlılığı (en yavaş / en hızlı dağılım süresi).", 12)
    SEKIL("57_tekrar_degiskenligi.png", "Beş tekrar arasındaki süre değişkenliği (varyasyon katsayısı).", 12)

    P(f"Şekil {SAYAC['sekil']-6}–{SAYAC['sekil']} dağılım ve tekrar farklarını ayırır. 100.000 "
      f"elemanda Insertion sıralı veride {uzun_sure(V['insertion','sirali',100000]['sure'])}, "
      f"ters sıralı veride {uzun_sure(V['insertion','ters_sirali',100000]['sure'])} sürer. "
      "Sıralı girdide kaydırma yapılmazken ters sırada her yeni eleman sol bölümün tamamını "
      "kaydırır. Beş rastgele tekrar farklı seed kullandığı için son grafikteki değişkenlik hem "
      "girdi farkını hem çalışma koşullarını içerir; yalnızca sistem gürültüsü olarak yorumlanamaz. "
      "Bubble ve Heap'in değişkenliği diğerlerinden yüksektir. Bu farkın sıcaklıkla ilişkisi "
      "kaydedilmediği için böyle bir neden ileri sürülmemiştir.")
    H("5.7 Bellek kullanımı", 2)
    bellek = {(s["algoritma"], int(s["boyut"])): s for s in oku(SONUCLAR / "bellek_sonuclari.csv")
              if s["veri_tipi"] == "rastgele" and s["basarili_mi"] == "True"}
    satirlar = []
    for a in ALGORITMALAR:
        s1 = bellek.get((a, 1000))
        s2 = bellek.get((a, 10000))
        satirlar.append([ALGO_AD[a],
                         tr(float(s1["olculen_tepe_bellek_bayt"])) if s1 else "–",
                         tr(float(s2["olculen_tepe_bellek_bayt"])) if s2 else "ölçülmedi",
                         (s1 or s2)["teorik_ek_alan"].replace("cagri yigini ust siniri", "çağrı yığını üst sınırı").replace("cagri yigini", "çağrı yığını")])
    TABLO("tracemalloc ile ölçülen tepe ek bellek (rastgele veri, bayt)",
          ["Algoritma", "n = 1.000", "n = 10.000", "Teorik ek alan"], satirlar,
          [3.0, 3.0, 3.2, 6.8], sagda=(1, 2), ilk_kalin=True)
    SEKIL("80_bellek_n10000.png", "Ek bellek kullanımı, n = 10.000 (log eksen).", 11.5)
    SEKIL("81_bellek_buyume.png", "Ek belleğin veri boyutuyla büyümesi.", 11.5)

    P(f"Şekil {SAYAC['sekil']-1}–{SAYAC['sekil']} için rastgele girdide Merge'in tepe tahsisi "
      f"{tr(float(bellek['merge',1000]['olculen_tepe_bellek_bayt']))} bayttan "
      f"{tr(float(bellek['merge',10000]['olculen_tepe_bellek_bayt']))} bayta çıkar. Kodda "
      "tek n elemanlı tampon bulunması bu artışı açıklar. Quick'in çağrı yığını ve Shell'in "
      "gap listesi daha yavaş büyür.")

    H("5.8 Bonus 1 – Quick Sort pivot karşılaştırması", 2)
    pivot = {(s["pivot_yontemi"], s["veri_tipi"], int(s["boyut"])): s
             for s in oku(SONUCLAR / "bonuslar" / "pivot_ozet.csv")}
    satirlar = []
    for d in DAGILIMLAR:
        o, i = pivot["orta", d, 10000], pivot["ilk", d, 10000]
        so, si = float(o["ortalama_sure_ms"]), float(i["ortalama_sure_ms"])
        satirlar.append([DAGILIM_AD[d], sure_tr(so), sure_tr(si), tr(si / so, 1) + "×",
                         tr(float(o["ortalama_karsilastirma"])), tr(float(i["ortalama_karsilastirma"]))])
    TABLO("Pivot yöntemleri, n = 10.000",
          ["Dağılım", "Orta (ms)", "İlk (ms)", "İlk / Orta", "Orta karşılaştırma", "İlk karşılaştırma"],
          satirlar, [3.1, 2.1, 2.2, 2.0, 3.3, 3.3], sagda=(1, 2, 3, 4, 5), ilk_kalin=True)
    SEKIL("60_pivot_sure.png", "Pivot yöntemlerinin süreye etkisi.", 15)
    SEKIL("61_pivot_oran.png", "İlk pivot süresinin orta pivot süresine oranı, n = 10.000.", 11.5)

    H("5.9 Bonus 2 – Shell Sort gap karşılaştırması", 2)
    gap = {(s["gap_yontemi"], s["veri_tipi"], int(s["boyut"])): s
           for s in oku(SONUCLAR / "bonuslar" / "gap_ozet.csv")}
    satirlar = []
    for d in DAGILIMLAR:
        k, nu = gap["klasik", d, 10000], gap["knuth", d, 10000]
        sk, sn = float(k["ortalama_sure_ms"]), float(nu["ortalama_sure_ms"])
        satirlar.append([DAGILIM_AD[d], sure_tr(sk), sure_tr(sn), tr(sk / sn, 2) + "×",
                         tr(float(k["ortalama_karsilastirma"])), tr(float(nu["ortalama_karsilastirma"]))])
    TABLO("Gap dizileri, n = 10.000",
          ["Dağılım", "Klasik (ms)", "Knuth (ms)", "Klasik / Knuth", "Klasik karşılaştırma", "Knuth karşılaştırma"],
          satirlar, [3.1, 2.1, 2.1, 2.4, 3.2, 3.1], sagda=(1, 2, 3, 4, 5), ilk_kalin=True)
    SEKIL("62_gap_sure.png", "Gap dizilerinin süreye etkisi.", 15)
    SEKIL("63_gap_oran.png", "Klasik gap süresinin Knuth gap süresine oranı, n = 10.000.", 11.5)

    H("5.10 Bonus 3 – n = 500.000", 2)
    buyuk = {(s["algoritma"], s["veri_tipi"]): s
             for s in oku(SONUCLAR / "bonuslar" / "buyuk_veri_ozet.csv")}
    algolar_buyuk = ["shell", "quick", "merge", "heap"]
    satirlar = []
    for d in DAGILIMLAR:
        satirlar.append([DAGILIM_AD[d]] + [uzun_sure(float(buyuk[a, d]["ortalama_sure_ms"])) for a in algolar_buyuk])
    TABLO("n = 500.000 için ortalama süre", ["Dağılım", "Shell", "Quick", "Merge", "Heap"],
          satirlar, [3.4, 3.1, 3.1, 3.1, 3.1], sagda=(1, 2, 3, 4), ilk_kalin=True)
    SEKIL("70_buyuk_veri_sure.png", "n = 500.000'de süre karşılaştırması.", 14)
    SEKIL("73_buyuk_veri_buyume.png", "n 100.000'den 500.000'e çıkınca sürenin büyüme çarpanı.", 11.5)

    P(f"Şekil {SAYAC['sekil']} rastgele girdide beş kat boyut artışının gerçek süre oranlarını "
      "gösterir: " + "; ".join(
          f"{ALGO_AD[a]} {tr(float(buyuk[a,'rastgele']['ortalama_sure_ms'])/V[a,'rastgele',100000]['sure'],2)} kat"
          for a in algolar_buyuk) + ". n·log n modeli bu aralıkta yaklaşık 5,70 kat artış verir. "
      "Model, aynı sabit maliyet varsayımı altında bir referanstır; oturum farklarını içermez. "
      "Shell'in oranı yalnızca gözlemdir ve n log n üst sınırı anlamına gelmez.")

    H("5.11 Bonus 4 – Paralel Merge Sort (teorik)", 2)
    P("Paralel Merge Sort yalnızca teorik olarak incelenmiştir; paralel kod yazılmamış, paralel ölçüm "
      "yapılmamıştır. Aşağıdaki iş ve kritik yol değerlendirmesi böl ve yönet modeline dayanır "
      "(Cormen ve diğerleri, 2022). Sıralı Merge Sort'ta sol ve sağ yarının sıralanması bağımsızdır ve eşzamanlı "
      "yürütülebilir; üst düzey birleştirme ise iki yarının tamamlanmasını bekler. Toplam iş "
      "T₁(n) = Θ(n log n); alt problemler paralel ama birleştirme sıralı olduğunda kritik yol "
      "T∞(n) = T∞(n/2) + Θ(n) = Θ(n) olur. Teorik paralellik T₁/T∞ = Θ(log n) ile sınırlıdır: çok "
      "sayıda çekirdek kullanılsa bile sıralı birleştirme darboğaz oluşturur. Amdahl yasası (Amdahl, 1967), paralel "
      "kesri f olan bir iş için S(p) = 1 / ((1 − f) + f/p) hızlanma verir; f = 0,9 için 2, 4 ve 8 "
      "çalışanla teorik hızlanma sırasıyla yaklaşık 1,82; 3,08 ve 4,71'dir. Bu değerler gerçek ölçüm "
      "değil, idealize bir model sonucudur. Ayrıca bu çalışmanın Python 3.13 ortamında GIL etkindir; "
      "saf Python thread'lerinin CPU ağırlıklı sıralamada paralel bayt kodu yürütmesi beklenmez, çok çekirdekli kazanç için "
      "ayrı süreçler kullanılabilir (Python Software Foundation, threading belgeleri). Süreç başlatma "
      "ve veri aktarım maliyetleri bu çalışmada ölçülmemiştir.")
    SEKIL("90_paralel_amdahl.png", "Amdahl yasasına göre teorik hızlanma (gerçek ölçüm değildir).", 11.5)
    SEKIL("91_paralel_sure_tahmini.png",
          "Ölçülen sıralı süre ve Amdahl modelinden hesaplanan teorik paralel süreler.", 12)

    H("5.12 Bonus 5 – Bubble Sort adım görselleştirmesi", 2)
    P("`gorsellestirme/bubble_gorsel.py` Tkinter ile yazılmış bağımsız bir eğitim uygulamasıdır. 5–20 tam "
      "sayıdan oluşan bir dizi girilir; **Sonraki**, **Önceki**, **Başlangıca Dön** ve **Sona Git** "
      "düğmeleriyle Bubble Sort'un karşılaştırma ve takas adımları sütunlar üzerinde izlenir. Uygulama "
      "performans ölçümü yapmaz. Adım mantığı birim testleriyle doğrulanmıştır.")
    P(f"Uygulama PyCharm üzerinden çalıştırıldı ve adım geçişleri tarafımdan kontrol "
      f"edildi. Şekil {SAYAC['sekil']+1} başlangıç dizisini, Şekil {SAYAC['sekil']+2} ilk takası, "
      f"Şekil {SAYAC['sekil']+3} tamamlanmış sıralamayı gösterir. Görseller kendi "
      "ekran görüntülerimden yalnızca kırpılarak alınmıştır; görüntü içeriği değiştirilmemiştir.")
    SEKIL(KOK / "rapor/gorseller/bubble_baslangic.png",
          "Bubble Sort başlangıcı: [8, 3, 5, 1, 4] dizisi, adım 0/18.", 15.5)
    SEKIL(KOK / "rapor/gorseller/bubble_takas.png",
          "İlk takas: 0 ve 1. indekslerdeki 8 ile 3 yer değiştirir; dizi [3, 8, 5, 1, 4] olur (adım 2/18).", 15.5)
    SEKIL(KOK / "rapor/gorseller/bubble_sonuc.png",
          "Tamamlanmış sıralama: [1, 3, 4, 5, 8]. Takassız son tur erken durmayı sağlar (adım 18/18).", 15.5)

    H("5.13 Tüm zorunlu deney sonuçları (ortalama, 5 tekrar)", 2)
    P("Aşağıdaki tablo tüm zorunlu kombinasyonları içerir; ham ölçümler (`ham_<n>.csv`) ve özetler "
      "(`ozet_<n>.csv`) ayrıca CSV olarak teslim edilmiştir.", boyut=12)
    satirlar = []
    for n in BOYUTLAR:
        for d in DAGILIMLAR:
            for a in ALGORITMALAR:
                r = V[a, d, n]
                satirlar.append([ALGO_AD[a], DAGILIM_AD[d], tr(n), sure_tr(r["sure"]),
                                 tr(r["karsilastirma"]), tr(r["hareket"])])
    TABLO("Tüm kombinasyonlar için ortalama süre, karşılaştırma ve hareket sayısı",
          ["Algoritma", "Veri tipi", "n", "Ort. süre (ms)", "Karşılaştırma", "Swap/Hareket"],
          satirlar, [2.5, 3.0, 1.8, 2.8, 3.0, 2.9], sagda=(2, 3, 4, 5), boyut=8.5)


    # =================================================== 6. KARŞILAŞTIRMA VE TARTIŞMA
    H("6. Karşılaştırma ve Tartışma")
    P("Aşağıda süre farklarını karşılaştırma, hareket ve veri düzeniyle birlikte değerlendiriyorum. "
      "Büyüme oranları aynı dağılımın farklı boyutlardaki ölçümlerinden hesaplanmıştır.")

    H("6.1 O(n²) algoritmaların süresi veri boyutu arttıkça nasıl değişti?", 2)
    sel = V["selection", "rastgele", 100000]["sure"] / V["selection", "rastgele", 1000]["sure"]
    ins = V["insertion", "rastgele", 100000]["sure"] / V["insertion", "rastgele", 1000]["sure"]
    bub = V["bubble", "rastgele", 100000]["sure"] / V["bubble", "rastgele", 1000]["sure"]
    P(f"Rastgele veride n 1.000'den 100.000'e (100 kat) çıkınca Selection Sort'un süresi "
      f"{tr(sel)} kat, Insertion Sort'un {tr(ins)} kat, Bubble Sort'un {tr(bub)} kat arttı. "
      "Karesel modelde n 100 kat arttığında yaklaşık 100² = 10.000 kat "
      "artış beklenir; oranlar bu referansla karşılaştırılmalıdır. n = 10.000'den 100.000'e deneysel "
      f"büyüme üssü (süre ∝ nᵏ) Selection için {tr(us('selection','rastgele'), 2)}, Insertion için "
      f"{tr(us('insertion','rastgele'), 2)}, Bubble için {tr(us('bubble','rastgele'), 2)} bulundu "
      "(aşağıdaki büyüme üssü grafiği). Karşılaştırma sayıları teoriyle tam örtüşür: Selection Sort her dağılımda tam olarak "
      f"n(n−1)/2 = {tr(V['selection','rastgele',100000]['karsilastirma'])} karşılaştırma yapar. "
      f"n = 100.000'de Selection Sort yaklaşık {uzun_sure(V['selection','rastgele',100000]['sure'])}, "
      f"Bubble Sort ise {uzun_sure(V['bubble','rastgele',100000]['sure'])} sürmüştür. Bubble Sort'un "
      "diğer ikisinden yavaş olmasının nedeni, çok sayıda takas yapmasıdır: rastgele veride Bubble "
      f"{tr(V['bubble','rastgele',100000]['hareket'])} veri hareketi yaparken Selection yalnızca "
      f"{tr(V['selection','rastgele',100000]['hareket'])} yapar.")
    SEKIL("51_buyume_ussu_rastgele.png", "Rastgele veride deneysel büyüme üssü (n: 10.000 → 100.000).", 12)
    SEKIL("50_normalize_sure_rastgele.png", "Normalize süre: süre/n² ve süre/(n·log₂ n), rastgele veri.", 15.5)

    H("6.2 O(n log n) algoritmalar neden büyük veri kümelerinde avantajlıdır?", 2)
    q = V["quick", "rastgele", 100000]
    P(f"100.000 elemanlı rastgele girdide Quick Sort {uzun_sure(q['sure'])} sürer. Selection/Quick "
      f"süre oranı {tr(V['selection','rastgele',100000]['sure']/q['sure'])}, Bubble/Quick oranı "
      f"{tr(V['bubble','rastgele',100000]['sure']/q['sure'])}'dür. Selection'ın her başlangıç "
      "konumunda kalan bölümün tamamını taraması, n(n−1)/2 karşılaştırmayı zorunlu kılar. Quick "
      "bu girdilerde bölümleri dengeli biçimde küçülterek daha az veri işlemi yapar.")
    P(f"Şekil {SAYAC['sekil']} süreyi n² ve n·log₂ n ile normalize eder. 100.000 elemanda "
      f"Merge'in {tr(V['merge','rastgele',100000]['karsilastirma'])} karşılaştırması, yaklaşık "
      "5 milyar karşılaştırma yapan Selection'dan çok düşüktür. Normalize eğrilerin ölçülen "
      "aralıkta yataylaşması ilgili büyüme modeliyle uyumludur. Shell aynı grafikte gösterilse de klasik gap için en kötü "
      "durum üst sınırı O(n²)'dir.")

    H("6.3 Insertion Sort sıralı verilerde neden başarılıdır?", 2)
    i = V["insertion", "sirali", 100000]
    P("Insertion'ın kaydırma sayısı girdideki ters çift sayısına eşittir. Sıralı girdide sol "
      "komşuyla yapılan ilk karşılaştırma döngüyü durdurur; buna rağmen geçiciye alma ve son "
      "yazma yapılır. 100.000 elemanda bu nedenle 99.999 karşılaştırma ve 199.998 hareket "
      f"kaydedilmiş, süre {uzun_sure(i['sure'])} olmuştur. Aynı veride erken duran Bubble "
      f"{uzun_sure(V['bubble','sirali',100000]['sure'])} sürmüş ve hiç hareket yapmamıştır.")
    P(f"Şekil {SAYAC['sekil']+1}'te tam sıralı ve yaklaşık %10 konumu değiştirilmiş veriyi ayırır. "
      "İkinci durumda uzak konumların değişmesi çok sayıda ters çift oluşturabilir: 100.000 "
      f"elemanda Insertion {tr(V['insertion','kismen_sirali',100000]['karsilastirma'])} "
      f"karşılaştırma yapmış ve {uzun_sure(V['insertion','kismen_sirali',100000]['sure'])} "
      "sürmüştür. Bu, elemanların çoğunun başlangıç konumunda kalmasının tek başına doğrusal "
      "süreyi garanti etmediğini gösterir. Ters sıralı girdide her yeni eleman sol bölümün "
      "tamamını kaydırır; maliyet O(n²)'ye ulaşır.")

    SEKIL("54_insertion_sirali_veri.png", "Insertion Sort sıralı ve kısmen sıralı veride (kalın çizgi).", 15)

    H("6.4 Quick Sort'un performansı pivot seçimine neden bağlıdır?", 2)
    oz = lambda p, d, a="ortalama_sure_ms": float(pivot[p, d, 10000][a])
    P(f"10.000 elemanlı sıralı girdide ilk pivot {tr(oz('ilk','sirali','ortalama_karsilastirma'))}, "
      f"orta pivot {tr(oz('orta','sirali','ortalama_karsilastirma'))} karşılaştırma yapmıştır. "
      f"İlk/orta süre oranı sıralı girdide {tr(oz('ilk','sirali')/oz('orta','sirali'),1)}, "
      f"ters sırada {tr(oz('ilk','ters_sirali')/oz('orta','ters_sirali'),1)}'tür. "
      "Bölüm 5.8'deki süre oranı grafiği bu farkı gösterir. Sıralı girdide ilk değer en küçük "
      "olduğu için bölümleme bir sonraki tura neredeyse bütün diziyi bırakır; toplam "
      "karşılaştırma sayısı karesel büyür. Orta indeks pivotu bu özel girdide dengeli bölünme sağlar.")
    P(f"Rastgele girdide aynı boyut için ilk/orta süre oranı "
      f"{tr(oz('ilk','rastgele')/oz('orta','rastgele'),2)}'dur. Bu, ilk pivotun her girdide aynı "
      "dezavantajı yaratmadığını gösterir. Ancak sabit bir orta indeks de bütün girdiler için "
      "iyi pivot garantisi vermez. Kod, küçük parçayı özyinelemeyle ve büyük parçayı döngüyle "
      "işlediğinden kötü bölümleme O(n²) zaman üretebilirken gerçek Python çağrı yığını "
      "O(log n) ile sınırlı kalır. Bu iki maliyet ayrı değerlendirilmelidir.")

    H("6.5 Merge Sort veri dağılımından nasıl etkileniyor?", 2)
    oran = {a: max(V[a, d, 100000]["sure"] for d in DAGILIMLAR) / min(V[a, d, 100000]["sure"] for d in DAGILIMLAR)
            for a in ALGORITMALAR}
    P("Merge Sort, `_araligi_sirala` içinde sınırları ortadan bölerek çalışır. Her birleştirme "
      "işlenen aralığı önce tampona, sonra ana diziye yazar; bu yüzden aynı boyuttaki hareket "
      "sayısı dağılımdan bağımsızdır. 100.000 elemanda dört dağılımda da "
      f"{tr(V['merge','rastgele',100000]['hareket'])} hareket kaydedilmiştir. Karşılaştırmalar "
      "ise yarılardan birinin ne zaman tükendiğine bağlıdır: sıralı girdide "
      f"{tr(V['merge','sirali',100000]['karsilastirma'])}, rastgele girdide "
      f"{tr(V['merge','rastgele',100000]['karsilastirma'])} karşılaştırma vardır.")
    P(f"En yavaş dağılım süresinin en hızlıya oranı Merge için {tr(oran['merge'],2)}, "
      f"Heap için {tr(oran['heap'],2)}, Selection için {tr(oran['selection'],2)} bulundu. "
      "Bu kayıtlar Merge'in dağılıma az duyarlı olduğunu gösterse de en düşük oran her zaman "
      "Merge'e ait değildir. Süre değişkenliği ile eşit anahtarların sırasının korunması "
      "anlamındaki kararlılık ayrı özelliklerdir. Merge'in kararlılığı, `_birlestir` içindeki "
      "`<=` koşulunda eşitlikte soldaki elemanın seçilmesinden gelir.")

    H("6.6 Heap Sort ile Merge Sort büyük veri kümelerinde nasıl karşılaştırılır?", 2)
    m, h = V["merge", "rastgele", 100000], V["heap", "rastgele", 100000]
    mb, hb = buyuk["merge", "rastgele"], buyuk["heap", "rastgele"]
    P(f"Rastgele 100.000 elemanda Merge {uzun_sure(m['sure'])}, Heap "
      f"{uzun_sure(h['sure'])} sürer; Heap/Merge süre oranı {tr(h['sure']/m['sure'],2)}'dir. "
      f"500.000 elemanda bu oran {tr(float(hb['ortalama_sure_ms'])/float(mb['ortalama_sure_ms']),2)} "
      "olur. Bunlar iki ayrı boyuttaki gerçek ortalamalardır; sabit bir evrensel hız oranı değildir.")
    P(f"100.000 elemanda Heap {tr(h['karsilastirma'])}, Merge {tr(m['karsilastirma'])} "
      f"karşılaştırma yapar. Hareket sayıları aynı sırayla {tr(h['hareket'])} ve "
      f"{tr(m['hareket'])}'dır. Heap'in `_asagi_suz` fonksiyonu çocukları tek tek karşılaştırıp "
      "gerektiğinde üç hareketlik takas yapar; Merge ise sıralı yarıları tek geçişte tampona "
      "aktarır. Sayaç farkları süre farkını açıklamaya yardımcı olur, ancak her Python işleminin "
      "aynı maliyette olduğunu göstermez.")
    P(f"Bellek deneyinde 10.000 eleman için Merge "
      f"{tr(float(bellek['merge',10000]['olculen_tepe_bellek_bayt']))}, Heap "
      f"{tr(float(bellek['heap',10000]['olculen_tepe_bellek_bayt']))} bayt tepe Python tahsisi "
      "üretmiştir. Kodda Merge'in n elemanlı tamponu, Heap'in ise yinelemeli ve tamponsuz "
      "süzmesi bulunur. Bu fark O(n) ve O(1) yardımcı alan açıklamalarıyla uyumludur. İşlemci "
      "önbelleği ayrıca ölçülmediğinden süre farkını kesin olarak önbelleğe bağlamıyorum.")

    H("6.7 Shell Sort'un performansı gap dizisinden nasıl etkilenir?", 2)
    gk = lambda y, d, a="ortalama_sure_ms": float(gap[y, d, 10000][a])
    P("Gap aralıkları, bir sonraki ekleme geçişine kalan işi değiştirir. Şekil "
      f"{SAYAC['sekil']+1}'te n=10.000 için rastgele veri karşılaştırma sayısı klasik yöntemde "
      f"{tr(gk('klasik','rastgele','ortalama_karsilastirma'))}, Knuth yönteminde "
      f"{tr(gk('knuth','rastgele','ortalama_karsilastirma'))}'tür. Ortalama süreler aynı sırayla "
      f"{sure_tr(gk('klasik','rastgele'))} ve {sure_tr(gk('knuth','rastgele'))} ms'dir. "
      "İki yöntem aynı kaydırma kodunu kullanır; fark gap listesinden gelir. Gap üretimi her iki "
      "algoritma çağrısının içinde olduğundan ölçülen süreye dahildir. Klasik yöntemin rastgele "
      f"veride 10.000–100.000 aralığındaki büyüme üssü {tr(us('shell','rastgele'),2)} bulundu. "
      "Bu değer yalnızca iki ölçüm boyutuna ait ampirik eğimdir; O(n log n) garantisi veya "
      "tüm girdiler için bir zaman karmaşıklığı sonucu değildir. Klasik yöntemin O(n²) en kötü "
      "durum sınırı geçerliliğini korur (Knuth, 1998).")

    SEKIL("62_gap_karsilastirma.png", "Klasik ve Knuth gap dizisi için karşılaştırma sayısı.", 15)

    H("6.8 Teorik Big-O değerleri deneysel sonuçlarla örtüşüyor mu?", 2)
    satirlar = []
    for a in ALGORITMALAR:
        satirlar.append([ALGO_AD[a], tr(us(a, "rastgele"), 2), tr(us(a, "sirali"), 2),
                         tr(us(a, "ters_sirali"), 2), tr(us(a, "kismen_sirali"), 2)])
    TABLO("Deneysel büyüme üssü k (süre ∝ nᵏ), n: 10.000 → 100.000",
          ["Algoritma", "Rastgele", "Sıralı", "Ters sıralı", "Kısmen sıralı"], satirlar,
          [3.6, 3.1, 3.1, 3.1, 3.1], sagda=(1, 2, 3, 4), ilk_kalin=True)
    P("Tablodaki üsler 10.000 ve 100.000 eleman için iki ortalama süreden hesaplanmıştır. "
      "Selection, Insertion ve Bubble'ın rastgele verideki üslerini 2 ile; Quick, Merge ve "
      "Heap'in değerlerini ise bu aralıkta n·log n modelinin verdiği yaklaşık 1,10 ile karşılaştırdım. "
      "Bu karşılaştırma yalnızca ölçülen aralığı açıklar. Algoritmanın bütün girdilerdeki davranışını "
      "kanıtlamak için süre eğrisi değil, kodun yaptığı işlemlerin analizi gerekir.")
    P(f"Kod ile kayıt arasında daha doğrudan bir kontrol Selection'da yapılabilir: 100.000 elemanda "
      f"{tr(V['selection','rastgele',100000]['karsilastirma'])} karşılaştırma, n(n−1)/2 ifadesiyle "
      "tam olarak aynıdır. Merge'in her birleştirmede tampona ve ana diziye yazması ise aynı "
      "boyutta dağılımdan bağımsız hareket sayısı üretir. Shell'in eğimi gap listesinden ve veri "
      "düzeninden etkilenir; tabloda n log n modeline yakın olması klasik yöntemin O(n²) üst "
      "sınırını değiştirmez.")
    MADDE(f"Insertion'ın sıralı girdideki üssü {tr(us('insertion','sirali'),2)} bulundu. Kodun "
          "n−1 karşılaştırma yaptığı kesindir; süre üssünün 1'den sapması ise farklı oturumlar, "
          "sabit maliyetler veya sistem gürültüsüyle ilişkili olabilir. Bu olasılıkların hangisinin "
          "etkili olduğunu belirleyecek ayrı bir ölçüm yapılmadı.")
    MADDE("Bellek tablosundaki bayt değerleri belirli Python tahsisleridir. Shell'in gap listesi "
          "ve Quick'in çağrı yığını için alan sınırı koddan O(log n) olarak çıkarılır; iki boyutta "
          "yakın tahsis değerleri ölçülmesi bu sınırı O(1)'e dönüştürmez.")

    H("6.9 Temel araştırma sorusunun yanıtı", 2)
    P("Algoritma tercihini boyut, dağılım ve ek alan gereksinimi birlikte belirler. 100.000 elemanlı "
      "rastgele veride bu kayıtların en hızlısı orta pivotlu Quick Sort'tur. Sıralı girdide ise "
      "Insertion ve erken duran Bubble yalnızca n−1 veri karşılaştırması yapar. Yaklaşık %10 "
      "konumu değiştirilen girdi, tam sıralı girdinin yerini tutmaz: uzak yer değiştirmeler çok "
      "sayıda kaydırma gerektirebilir.")
    P("En kötü durumda O(n log n) zaman sınırı isteniyorsa Merge ve Heap bu uygulamalar içinde "
      "uygun seçeneklerdir. Merge ek tampon karşılığında kararlı sıralama sağlar; Heap yinelemeli "
      "süzme sayesinde sabit ek alan kullanır. Shell'in klasik ve Knuth sonuçları gap seçiminin "
      "önemini gösterir, ancak ölçülen hız avantajı bütün veri kümelerine genellenemez.")


    # ===================================================================== 7. SONUÇ
    H("7. Sonuç")
    P("Zorunlu deney matrisindeki 840 ölçümün tamamında sıralılık ve eleman bütünlüğü sağlandı. "
      "Deneyler, büyük rastgele girdilerde yapılan veri işlemlerinin süre farkını açıklamada yararlı "
      "olduğunu; sıralı veride ise erken durma ve kaydırma ihtiyacının sonucu değiştirdiğini gösterdi.")
    P(f"100.000 elemanlı rastgele veride Quick Sort {uzun_sure(q['sure'])}, Merge Sort "
      f"{uzun_sure(m['sure'])}, Heap Sort {uzun_sure(h['sure'])} sürdü. İlk pivotun sıralı girdide "
      f"orta pivota göre süre oranı, 10.000 elemanda {tr(oz('ilk','sirali')/oz('orta','sirali'),1)} "
      "olarak ölçüldü. Bu fark çağrı yığınının büyümesinden değil, dengesiz bölünmelerde artan "
      "veri karşılaştırmalarından kaynaklanır. Shell için gözlenen eğim gap seçimine bağlıdır ve "
      "klasik yöntemin O(n²) üst sınırıyla karıştırılmamalıdır.")
    P("Sonuçlar bu Python gerçeklemeleri ve kayıtlı deney koşullarıyla sınırlıdır. Donanım sıcaklığı "
      "izlenmediği için süre farklarını ısınmaya bağlayamam. Paralel Merge için verilen sayılar model "
      "tahminidir; gerçek paralel hızlanma ölçülmemiştir. Bu ayrımlar, hangi sonuçların doğrudan "
      "ölçüme dayandığını ve hangilerinin ayrıca sınanması gerektiğini belirler.")


    # =================================================================== 8. KAYNAKÇA
    H("8. Kaynakça")
    for kaynak in (
        "Cormen, T. H., Leiserson, C. E., Rivest, R. L., & Stein, C. (2022). Introduction to Algorithms "
        "(4th ed.). MIT Press.",
        "Sedgewick, R., & Wayne, K. (2011). Algorithms (4th ed.). Addison-Wesley.",
        "Knuth, D. E. (1998). The Art of Computer Programming, Volume 3: Sorting and Searching (2nd ed.). "
        "Addison-Wesley.",
        "Hoare, C. A. R. (1962). Quicksort. The Computer Journal, 5(1), 10–16.",
        "Amdahl, G. M. (1967). Validity of the single processor approach to achieving large scale computing "
        "capabilities. AFIPS Spring Joint Computer Conference, 483–485.",
        "Python Software Foundation. Python 3.13 documentation: time.perf_counter_ns, tracemalloc, "
        "threading. https://docs.python.org/3.13/library/time.html ; "
        "https://docs.python.org/3.13/library/tracemalloc.html ; "
        "https://docs.python.org/3.13/library/threading.html (erişim: 23.10.2026).",
        "Sedgewick, R., & Wayne, K. Shell sınıfı belgeleri. "
        "https://algs4.cs.princeton.edu/code/javadoc/edu/princeton/cs/algs4/Shell.html (erişim: 23.10.2026).",
    ):
        paragraf = belge.add_paragraph()
        paragraf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _calistir(paragraf, kaynak)
    P("Proje Kaynak Koduna Erişim", kalin=True, hiza=WD_ALIGN_PARAGRAPH.LEFT)
    P("Kaynak kod, otomatik testler, kayıtlı deney sonuçları ve README için proje deposu: "
      "https://github.com/Gulced/AlgoritmaSiralamaKarsilastirma")

    # Sayfa numaraları görüntüleyicide hesaplanır; içindekiler ana bölüm listesidir.
    ayarlar = belge.settings.element
    guncelle = OxmlElement("w:updateFields")
    guncelle.set(qn("w:val"), "true")
    ayarlar.append(guncelle)

    belge.save(cikti)
    print("Yazıldı:", cikti)
    print("Şekil:", SAYAC["sekil"], "Tablo:", SAYAC["tablo"])

    return belge


if __name__ == "__main__":
    import argparse
    import subprocess
    secenekler = argparse.ArgumentParser()
    secenekler.add_argument("--word-alanlarini-guncelle", action="store_true",
                           help="macOS Word ile içindekiler ve sayfa numaralarını günceller; PDF üretmez.")
    secenekler.add_argument("--pdf", action="store_true", help="Word alanlarını güncelleyip teslim PDF'sini de üretir.")
    args = secenekler.parse_args()
    raporOlustur()
    if args.word_alanlarini_guncelle or args.pdf:
        subprocess.run(["/usr/bin/osascript", str(KOK / "rapor/word_alanlarini_guncelle.applescript"),
                        str(CIKTI)], check=True, timeout=120)
    if args.pdf:
        import shutil
        soffice = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"
        if not soffice.exists():
            soffice = Path(shutil.which("soffice") or "soffice")
        hedef = KOK / "output/pdf"
        hedef.mkdir(parents=True, exist_ok=True)
        import tempfile
        with tempfile.TemporaryDirectory(prefix="bmt317-pdf-") as profil:
            subprocess.run([str(soffice), "-env:UserInstallation=" + Path(profil).as_uri(), "--headless",
                            "--convert-to", "pdf", "--outdir", str(hedef), str(CIKTI)], check=True, timeout=120)
